from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUTPUT = ROOT / "BaoCao_DoAn.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def add_toc(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Mục lục sẽ được cập nhật khi mở tài liệu trong Microsoft Word."
    result_run = OxmlElement("w:r")
    result_run.append(placeholder)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, result_run, end])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(5)

    heading_specs = {
        "Title": (20, "17365D", True),
        "Heading 1": (16, "17365D", True),
        "Heading 2": (14, "1F4E78", True),
        "Heading 3": (13, "365F91", True),
        "Heading 4": (12, "365F91", True),
    }
    for name, (size, color, bold) in heading_specs.items():
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        # Large screenshots and tables may follow a heading. Let Word paginate
        # them independently so an oversized heading+object group never spills
        # into the physical top margin.
        style.paragraph_format.keep_with_next = False

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(12)

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = "Consolas"
        code._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        code.font.size = Pt(8.5)
        code.paragraph_format.left_indent = Cm(0.5)
        code.paragraph_format.right_indent = Cm(0.3)
        code.paragraph_format.space_before = Pt(3)
        code.paragraph_format.space_after = Pt(3)

    if "Caption Custom" not in styles:
        cap = styles.add_style("Caption Custom", WD_STYLE_TYPE.PARAGRAPH)
        cap.font.name = "Times New Roman"
        cap._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        cap.font.size = Pt(10.5)
        cap.font.italic = True
        cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(8)

    header = section.header.paragraphs[0]
    header.text = "NHÓM 9 · HỆ THỐNG QUẢN LÝ ĐỀ TÀI SINH VIÊN"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.name = "Times New Roman"
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor(100, 100, 100)
    add_page_number(section.footer.paragraphs[0])


def add_inline(paragraph, text: str) -> None:
    text = text.replace("$\\to$", "→").replace("$→$", "→")
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\([^)]+\)|\*[^*]+\*)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(192, 57, 43)
        elif token.startswith("["):
            label = token[1:token.index("]")]
            url = token[token.index("(") + 1:-1]
            run = paragraph.add_run(f"{label} ({url})")
            run.font.color.rgb = RGBColor(31, 78, 121)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, row_obj in enumerate(table.rows):
        tr_pr = row_obj._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        if row_index == 0:
            table_header = OxmlElement("w:tblHeader")
            table_header.set(qn("w:val"), "true")
            tr_pr.append(table_header)
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            text = row[j].strip() if j < len(row) else ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            add_inline(paragraph, text)
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(9.5)
                if i == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            if i == 0:
                set_cell_shading(cell, "1F4E78")
            elif i % 2 == 0:
                set_cell_shading(cell, "EAF2F8")
    doc.add_paragraph()


def resolve_image(source_md: Path, raw: str) -> Path | None:
    candidate = (source_md.parent / raw).resolve()
    if candidate.suffix.lower() == ".svg":
        png = candidate.with_suffix(".png")
        if png.exists():
            candidate = png
    return candidate if candidate.exists() else None


def add_image(doc: Document, path: Path, alt: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    try:
        run = p.add_run()
        run.add_picture(str(path), width=Inches(6.25))
    except Exception:
        add_inline(p, f"[{alt}: không thể nhúng ảnh {path.name}]")


def add_markdown(doc: Document, source: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    i = 0
    in_code = False
    code_lines: list[str] = []
    skipped_wrapper = False

    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                p = doc.add_paragraph(style="Code Block")
                p.paragraph_format.keep_together = True
                p.add_run("\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        image_match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            path = resolve_image(source, image_match.group(2))
            if path:
                add_image(doc, path, image_match.group(1))
            i += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows: list[list[str]] = []
            for idx, table_line in enumerate(table_lines):
                cells = [c.strip() for c in table_line.strip("|").split("|")]
                if idx == 1 and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells):
                    continue
                rows.append(cells)
            add_table(doc, rows)
            continue

        heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if source.name == "chapter-1-2.md" and not skipped_wrapper and text.startswith("Nội dung báo cáo"):
                skipped_wrapper = True
                i += 1
                continue
            if source.name == "chapter-4.md" and text.startswith("4.4 "):
                doc.add_page_break()
            if source.name == "chapter-5-6.md" and re.match(r"5\.[2-9]\.", text):
                doc.add_page_break()
            if source.name == "chapter-5-6.md" and text.startswith("6.4."):
                doc.add_page_break()
            if re.match(r"Chương\s+[1-6]", text, re.IGNORECASE):
                if len(doc.paragraphs) > 5:
                    doc.add_page_break()
                p = doc.add_paragraph(style="Heading 1")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(p, text.replace(":", ".", 1) if text.startswith("Chương 3:") else text)
            else:
                p = doc.add_paragraph(style=f"Heading {min(level, 4)}")
                add_inline(p, text)
            i += 1
            continue

        if not stripped or stripped == "---":
            i += 1
            continue

        if re.fullmatch(r"\*Hình\s+.+\*", stripped):
            p = doc.add_paragraph(style="Caption Custom")
            p.add_run(stripped[1:-1])
            i += 1
            continue

        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.right_indent = Cm(0.4)
            p.paragraph_format.keep_together = True
            add_inline(p, stripped.lstrip("> "))
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(80, 80, 80)
            i += 1
            continue

        bullet = re.match(r"^[-*]\s+(.+)$", stripped)
        number = re.match(r"^\d+[.)]\s+(.+)$", stripped)
        if bullet or number:
            if bullet:
                p = doc.add_paragraph(style="List Bullet")
                add_inline(p, bullet.group(1))
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.55)
                p.paragraph_format.first_line_indent = Cm(-0.55)
                add_inline(p, f"{stripped.split(maxsplit=1)[0]} {number.group(1)}")
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline(p, stripped)
        i += 1


def add_cover(doc: Document) -> None:
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo = ROOT / "src/main/resources/static/images/hcmute-logo.jpg"
    if logo.exists():
        p.add_run().add_picture(str(logo), width=Inches(1.05))

    for text, size, bold, color, space in [
        ("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ KỸ THUẬT TP. HỒ CHÍ MINH", 13, True, "17365D", 0),
        ("KHOA CÔNG NGHỆ THÔNG TIN", 14, True, "17365D", 20),
        ("BÁO CÁO ĐỒ ÁN MÔN HỌC", 20, True, "17365D", 8),
        ("HỆ THỐNG QUẢN LÝ ĐỀ TÀI SINH VIÊN", 22, True, "C00000", 22),
        ("NHÓM 9", 16, True, "1F4E78", 10),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(space)
        r = p.add_run(text)
        r.bold = bold
        r.font.name = "Times New Roman"
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor.from_string(color)

    table = doc.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    data = [
        ["Thành viên", "MSSV", "Phụ trách"],
        ["Trang Sĩ Hoàng", "24162035", "Core, Admin, tích hợp"],
        ["Vũ Trọng Hưng", "24162053", "Đề tài, nhóm, đăng ký, báo cáo"],
        ["Trần Hào Kiệt", "24162065", "Phản biện, hội đồng, điểm, công bố, thông báo, dashboard, triển khai"],
    ]
    for i, row in enumerate(data):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            set_cell_margins(cell, 110, 100, 110, 100)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(value)
            r.font.name = "Times New Roman"
            r.font.size = Pt(10.5)
            r.bold = i == 0
            if i == 0:
                r.font.color.rgb = RGBColor(255, 255, 255)
                set_cell_shading(cell, "1F4E78")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    r = p.add_run("TP. HỒ CHÍ MINH, THÁNG 09 NĂM 2026")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    doc.add_page_break()


def add_front_matter(doc: Document) -> None:
    p = doc.add_paragraph("LỜI CAM ĐOAN", style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    text = (
        "Nhóm 9 cam đoan báo cáo phản ánh đúng sản phẩm đã xây dựng và kết quả kiểm thử thực tế. "
        "Các ảnh giao diện được chụp trực tiếp từ ứng dụng; thông tin mật, mật khẩu và token không được đưa vào báo cáo hoặc gói nộp."
    )
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    p = doc.add_paragraph("MỤC LỤC", style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc = doc.add_paragraph()
    add_toc(toc)


def add_test_appendix(doc: Document) -> None:
    p = doc.add_paragraph("PHỤ LỤC. KẾT QUẢ KIỂM THỬ", style="Heading 1")
    p.paragraph_format.page_break_before = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "Ngày 24/09/2026, hệ thống được kiểm thử đủ 16 bước trên MySQL 8.0.46 bằng Maven tại cổng 8090. "
        "Kết quả cuối: 44/44 test tự động đạt; 4 phiếu chấm bị khóa; điểm tổng kết 8,79; Chủ tịch xác nhận; Khoa công bố; "
        "sinh viên trong nhóm xem được và sinh viên ngoài nhóm nhận 403."
    )
    rows = [
        ["Hạng mục", "Kết quả"],
        ["JUnit", "44 test, 0 failure, 0 error, 0 skipped"],
        ["MySQL / Flyway", "MySQL 8.0.46; schema version 10"],
        ["Health / static", "HTTP 200, UP; CSS HTTP 200"],
        ["E2E", "16/16 bước đạt trên MySQL/Maven"],
        ["Docker", "Chưa chạy lại vì máy không có Docker Desktop/CLI"],
        ["Production", "Chưa có URL do thiếu tài khoản/credential hosting"],
    ]
    add_table(doc, rows)


def main() -> None:
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_front_matter(doc)
    for filename in ("chapter-1-2.md", "chapter-3.md", "chapter-4.md", "chapter-5-6.md"):
        add_markdown(doc, DOCS / filename)
    add_test_appendix(doc)

    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.widow_control = True
        for run in paragraph.runs:
            if not run.font.name:
                run.font.name = "Times New Roman"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    doc.core_properties.title = "Báo cáo đồ án – Hệ thống quản lý đề tài sinh viên"
    doc.core_properties.subject = "Báo cáo Chương 1–6 và kết quả kiểm thử E2E"
    doc.core_properties.author = "Nhóm 9 – 24162035, 24162053, 24162065"
    doc.core_properties.keywords = "Spring Boot, MySQL, quản lý đề tài, Nhóm 9"
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    doc.settings._element.append(update_fields)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
